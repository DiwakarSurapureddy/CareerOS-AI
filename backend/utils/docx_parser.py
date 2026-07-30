import logging
from docx import Document

logger = logging.getLogger(__name__)

def extract_text_from_docx(file_path: str):
    """
    Extract text from paragraphs and tables of a DOCX file using python-docx.
    Returns (True, extracted_text) on success, or (False, error_message) on failure.
    """
    try:
        doc = Document(file_path)
    except Exception as e:
        logger.error(f"Failed to load DOCX file at {file_path}: {e}")
        return False, "Failed to open or parse the DOCX file. The file may be invalid or corrupted."
        
    try:
        text_parts = []
        
        # Extract paragraph text
        for para in doc.paragraphs:
            if para.text and para.text.strip():
                text_parts.append(para.text.strip())
                
        # Extract text from tables if available
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    cell_txt = cell.text.strip()
                    if cell_txt and cell_txt not in row_text:
                        row_text.append(cell_txt)
                if row_text:
                    text_parts.append(" | ".join(row_text))
                    
        combined_text = "\n\n".join(text_parts).strip()
        if not combined_text:
            return False, "The DOCX document appears to be empty or contains no readable text."
            
        return True, combined_text
    except Exception as e:
        logger.error(f"Error extracting text from DOCX {file_path}: {e}")
        return False, f"Error extracting text from DOCX file: {str(e)}"
